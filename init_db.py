import os
import datetime
import openpyxl
from database import engine, SessionLocal, Base
from models import User, BudgetCategory, TreasuryAccount, Transaction, AccountMonthlyBalance
from auth import hash_password

def init_database():
    print("Creating database tables...")
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()

    # 1. Crear usuarios por defecto
    users_data = [
        ("directivo", "tev2026*", "Dirección General TEV", "directivo"),
        ("administradora", "admin2026*", "Administradora TEV", "administradora"),
        ("cajera1", "caja12026*", "Cajera / Asistente 1", "cajera"),
        ("cajera2", "caja22026*", "Cajera / Asistente 2", "cajera"),
    ]

    for username, pwd, name, role in users_data:
        existing = db.query(User).filter(User.username == username).first()
        if not existing:
            user = User(
                username=username,
                password_hash=hash_password(pwd),
                full_name=name,
                role=role,
                is_active=True
            )
            db.add(user)
            print(f"User created: {username} ({role})")
    db.commit()

    admin_user = db.query(User).filter(User.username == "administradora").first()

    # 2. Cuentas de Tesorería por defecto
    accounts_data = [
        ("Efectivo USD", "USD", "EFECTIVO", 0.0, False),
        ("Efectivo VES", "VES", "EFECTIVO", 0.0, False),
        ("Banco Banesco VES", "VES", "BANCO", 0.0, False),
        ("Banco Mercantil VES", "VES", "BANCO", 0.0, False),
        ("BNC - CASHEA", "VES", "BANCO", 0.0, True),
        ("Banco Bancaribe VES", "VES", "BANCO", 0.0, False),
        ("Banco de Venezuela VES", "VES", "BANCO", 0.0, False),
        ("Zelle USD", "USD", "BANCO", 0.0, False),
        ("Billetera USDT", "USDT", "BILLETERA", 0.0, False),
    ]

    for name, curr, acc_type, init_bal, inc_only in accounts_data:
        existing = db.query(TreasuryAccount).filter(TreasuryAccount.name == name).first()
        if not existing:
            acc = TreasuryAccount(
                name=name,
                currency=curr,
                account_type=acc_type,
                initial_balance=init_bal,
                only_income=inc_only,
                is_active=True
            )
            db.add(acc)
            print(f"Account created: {name} [{curr}]")
    db.commit()

    # Configuración de tasa BCV inicial
    from models import SystemSetting
    bcv = db.query(SystemSetting).filter(SystemSetting.key == 'bcv_rate').first()
    if not bcv:
        db.add(SystemSetting(key='bcv_rate', value='36.80', updated_by='sistema'))
        db.commit()

    acc_usd = db.query(TreasuryAccount).filter(TreasuryAccount.name == "Efectivo USD").first()
    acc_ves = db.query(TreasuryAccount).filter(TreasuryAccount.name == "Banco Banesco VES").first()

    # 3. Cargar las 12 Partidas Presupuestarias
    categories_budget = [
        (1, "Impuestos Seniat", 1500.0),
        (2, "Impuestos Municipales", 500.0),
        (3, "Parafiscales", 100.0),
        (4, "Servicios", 1000.0),
        (5, "Gastos Operativos y Mantenimiento", 500.0),
        (6, "Alquileres", 2250.0),
        (7, "Nomina y Pasivos Laborales", 4500.0),
        (8, "Comisiones por venta", 1200.0),
        (9, "Gastos Extraordinarios", 1000.0),
        (10, "Compras de Bienes", 500.0),
        (11, "Retiros de Accionista", 750.0),
        (12, "Mantenimiento Flota", 200.0),
    ]

    cat_map = {}
    for code, name, budget in categories_budget:
        existing = db.query(BudgetCategory).filter(BudgetCategory.code == code).first()
        if not existing:
            cat = BudgetCategory(code=code, name=name, monthly_budget_usd=budget, is_active=True)
            db.add(cat)
            db.flush()
            cat_map[code] = cat.id
            print(f"Category created: {code} - {name} (${budget})")
        else:
            cat_map[code] = existing.id
    db.commit()

    # 4. Importar Histórico de Gastos de Agosto 2026
    excel_path = "C:/Users/GATEWAY/Desktop/CLIENTES DE CONSULTORIA/TODO ELECTRICO VALENCIA/ARCHIVOS DE SEGUIMIENTO 2026/CONTROL DE GASTOS 08 AGOSTO.xlsx"
    if os.path.exists(excel_path):
        wb = openpyxl.load_workbook(excel_path, data_only=True)
        if "Gastos Ordinarios Mes a Mes" in wb.sheetnames:
            ws = wb["Gastos Ordinarios Mes a Mes"]
            count_trans = db.query(Transaction).count()
            if count_trans == 0:
                print("Importing August 2026 expenses from Excel...")
                imported = 0
                for r in range(2, ws.max_row + 1):
                    raw_date = ws.cell(r, 1).value
                    desc = ws.cell(r, 3).value
                    cat_code = ws.cell(r, 4).value
                    monto_bs = ws.cell(r, 6).value
                    monto_usd = ws.cell(r, 7).value
                    tasa = ws.cell(r, 8).value

                    if not raw_date or not desc:
                        continue

                    if isinstance(raw_date, datetime.datetime):
                        t_date = raw_date.date()
                    elif isinstance(raw_date, datetime.date):
                        t_date = raw_date
                    else:
                        try:
                            t_date = datetime.datetime.strptime(str(raw_date)[:10], "%Y-%m-%d").date()
                        except:
                            t_date = datetime.date(2026, 8, 1)

                    try:
                        monto_bs_flt = float(monto_bs) if monto_bs is not None else 0.0
                    except:
                        monto_bs_flt = 0.0

                    try:
                        monto_usd_flt = float(monto_usd) if monto_usd is not None else 0.0
                    except:
                        monto_usd_flt = 0.0

                    try:
                        tasa_flt = float(tasa) if tasa is not None else 756.0
                    except:
                        tasa_flt = 756.0

                    if monto_bs_flt > 0:
                        account_id = acc_ves.id
                        currency = "VES"
                        amount_orig = monto_bs_flt
                    else:
                        account_id = acc_usd.id
                        currency = "USD"
                        amount_orig = monto_usd_flt

                    cat_id = None
                    try:
                        code_int = int(cat_code)
                        cat_id = cat_map.get(code_int)
                    except:
                        cat_id = cat_map.get(9)

                    tx = Transaction(
                        date=t_date,
                        movement_type="EGRESO",
                        subtype="GASTO_OPERATIVO",
                        account_id=account_id,
                        category_id=cat_id,
                        amount_original=amount_orig,
                        currency=currency,
                        exchange_rate=tasa_flt,
                        amount_usd=monto_usd_flt,
                        reference_number="",
                        beneficiary="",
                        description=str(desc).strip(),
                        status="VERIFICADO",
                        created_by_id=admin_user.id,
                        verified_by_id=admin_user.id
                    )
                    db.add(tx)
                    imported += 1

                db.commit()
                print(f"Imported {imported} transactions from August 2026 successfully.")

    db.close()
    print("Database initialization complete.")

if __name__ == "__main__":
    init_database()
