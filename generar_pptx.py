import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 widescreen format
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette TEV
    NAVY = RGBColor(15, 43, 72)        # #0f2b48
    GOLD = RGBColor(245, 158, 11)      # #f59e0b
    DARK_BLUE = RGBColor(26, 54, 93)   # #1a365d
    LIGHT_BG = RGBColor(248, 250, 252) # #f8fafc
    WHITE = RGBColor(255, 255, 255)
    GRAY_TEXT = RGBColor(100, 116, 139)
    DARK_TEXT = RGBColor(30, 41, 59)
    CARD_BG = RGBColor(255, 255, 255)
    CARD_BORDER = RGBColor(226, 232, 240)

    logo_path = os.path.abspath("static/logo.jpg")

    def add_header(slide, title_text, category_text="TODO ELÉCTRICO VALENCIA, C.A. • CONSULTORÍA DE GESTIÓN"):
        # Header banner shape
        banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.15))
        banner.fill.solid()
        banner.fill.fore_color.rgb = NAVY
        banner.line.color.rgb = NAVY

        # Gold accent line
        accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.15), Inches(13.333), Inches(0.08))
        accent.fill.solid()
        accent.fill.fore_color.rgb = GOLD
        accent.line.color.rgb = GOLD

        # Category text
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.18), Inches(9.5), Inches(0.3))
        cat_tf = cat_box.text_frame
        cat_tf.word_wrap = True
        p_cat = cat_tf.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = GOLD

        # Title text
        t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(9.5), Inches(0.6))
        t_tf = t_box.text_frame
        t_tf.word_wrap = True
        p_t = t_tf.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(22)
        p_t.font.bold = True
        p_t.font.color.rgb = WHITE

        # Logo on right
        if os.path.exists(logo_path):
            slide.shapes.add_picture(logo_path, Inches(11.2), Inches(0.15), height=Inches(0.85))

        # Bottom footer
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.1), Inches(11.733), Inches(0.3))
        ft_tf = footer_box.text_frame
        p_ft = ft_tf.paragraphs[0]
        p_ft.text = "Sistema de Tesorería y Flujo de Caja • Confidencial para Todo Eléctrico Valencia, C.A."
        p_ft.font.size = Pt(9)
        p_ft.font.color.rgb = GRAY_TEXT

    blank_layout = prs.slide_layouts[6]

    # ==========================================
    # SLIDE 1: PORTADA
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = NAVY
    bg1.line.color.rgb = NAVY

    # Decorative Gold Bar
    bar1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(1.6), Inches(0.15), Inches(4.2))
    bar1.fill.solid()
    bar1.fill.fore_color.rgb = GOLD
    bar1.line.color.rgb = GOLD

    # Logo
    if os.path.exists(logo_path):
        s1.shapes.add_picture(logo_path, Inches(1.6), Inches(1.5), height=Inches(1.2))

    # Title & Subtitle Box
    tbox1 = s1.shapes.add_textbox(Inches(1.6), Inches(2.9), Inches(10.5), Inches(3.2))
    tf1 = tbox1.text_frame
    tf1.word_wrap = True

    p1 = tf1.paragraphs[0]
    p1.text = "Control Inteligente de Tesorería y Flujo de Caja"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = WHITE

    p2 = tf1.add_paragraph()
    p2.text = "Evolución Estratégica: De la Gestión Convencional a la Automatización Financiera en Tiempo Real"
    p2.font.size = Pt(18)
    p2.font.color.rgb = GOLD
    p2.space_before = Pt(12)

    p3 = tf1.add_paragraph()
    p3.text = "Todo Eléctrico Valencia, C.A. • RIF: J-50248654-2\nAv. Monseñor Adams, El Viñedo. Valencia, Carabobo. Venezuela."
    p3.font.size = Pt(13)
    p3.font.color.rgb = RGBColor(203, 213, 225)
    p3.space_before = Pt(18)

    # ==========================================
    # SLIDE 2: OBJETIVO DEL ENCUENTRO
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "¿Por qué estamos hoy aquí?", "PROPÓSITO Y ENFOQUE ESTRATÉGICO")
    
    # 3 Cards
    cards_data_s2 = [
        ("1. Claridad Financiera", "Comprender la salud real del negocio sin tecnicismos contables ni hojas complejas. Ver el dinero tal como fluye día a día.", GOLD),
        ("2. Tranquilidad Operativa", "Eliminar la incertidumbre de cuadres nocturnos y la dispersión de pagos entre Bolívares, Dólares en efectivo y Zelle.", DARK_BLUE),
        ("3. Decisión Oportuna", "Empoderar a la junta directiva y administración con números exactos a un solo clic para anticipar pagos y compras.", NAVY)
    ]
    for i, (ctitle, cdesc, ccolor) in enumerate(cards_data_s2):
        left = Inches(1.0 + i * 3.9)
        top = Inches(2.0)
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(3.6), Inches(4.5))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1.5)

        # Top tag
        tag = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(3.6), Inches(0.7))
        tag.fill.solid()
        tag.fill.fore_color.rgb = ccolor
        tag.line.color.rgb = ccolor
        
        tf_tag = tag.text_frame
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = ctitle
        p_tag.font.size = Pt(15)
        p_tag.font.bold = True
        p_tag.font.color.rgb = WHITE
        p_tag.alignment = PP_ALIGN.CENTER

        # Body text
        tb = s2.shapes.add_textbox(left + Inches(0.3), top + Inches(1.1), Inches(3.0), Inches(3.0))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = cdesc
        p.font.size = Pt(14)
        p.font.color.rgb = DARK_TEXT

    # ==========================================
    # SLIDE 3: ¿QUÉ ES EL FLUJO DE CAJA?
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "¿Qué es el Flujo de Caja? (El Oxígeno del Negocio)", "CONCEPTO FINANCIERO CLAVE 1")

    # Left Column: Analogy
    box_left = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.8), Inches(5.4), Inches(4.8))
    box_left.fill.solid()
    box_left.fill.fore_color.rgb = WHITE
    box_left.line.color.rgb = GOLD
    box_left.line.width = Pt(2)

    tf_l = box_left.text_frame
    tf_l.word_wrap = True
    p = tf_l.paragraphs[0]
    p.text = "La Analogía del Tanque de Agua"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY

    p_body = tf_l.add_paragraph()
    p_body.text = (
        "\nImagine que su empresa es un tanque de agua:\n\n"
        "• Entradas de Agua (Ingresos): Las ventas cobradas hoy, transferencias recibidas y abonos de clientes.\n\n"
        "• Salidas de Agua (Egresos): Compras a proveedores, pago de luz, nóminas, fletes y viáticos.\n\n"
        "• Nivel Actual (Saldo Disponible): El agua que realmente le queda en el tanque para seguir operando mañana."
    )
    p_body.font.size = Pt(13)
    p_body.font.color.rgb = DARK_TEXT

    # Right Column: Key Distinction
    box_right = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.5), Inches(4.8))
    box_right.fill.solid()
    box_right.fill.fore_color.rgb = NAVY
    box_right.line.color.rgb = NAVY

    tf_r = box_right.text_frame
    tf_r.word_wrap = True
    p = tf_r.paragraphs[0]
    p.text = "La Regla de Oro en Todo Negocio"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = GOLD

    p_body_r = tf_r.add_paragraph()
    p_body_r.text = (
        "\n\"Una empresa no quiebra por falta de ventas;\nquiebra por falta de liquidez en caja.\"\n\n"
        "Usted puede facturar $100.000 este mes en el papel contable, pero si la mitad se vendió a crédito o el dinero está atrapado en inventario, mañana no tendrá con qué pagar la nómina ni la reposición urgente de material.\n\n"
        "El Flujo de Caja no le habla del futuro prometido, le dice la verdad estricta de hoy."
    )
    p_body_r.font.size = Pt(13)
    p_body_r.font.color.rgb = WHITE

    # ==========================================
    # SLIDE 4: CONTROL DE GASTOS ESTRATÉGICO
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "¿Qué es un Control de Gastos Inteligente?", "CONCEPTO FINANCIERO CLAVE 2")

    items_s4 = [
        ("No se trata de no gastar, sino de gastar con sentido", "El objetivo no es frenar el negocio, sino asegurar que cada dólar o bolívar gastado genere valor o mantenga la operatividad activa.", DARK_BLUE),
        ("Clasificación Automática por Rubros", "Saber con precisión matemática cuánto se destina a: Operaciones, Reposición de Mercancía, Nómina, Mantenimiento e Impuestos.", NAVY),
        ("Detección Inmediata de Fugas Invisibles", "Los pequeños gastos dispersos de caja chica (almuerzos, traslados, compras imprevistas) que sin registro suman miles de dólares al año.", GOLD)
    ]
    for i, (ititle, idesc, icolor) in enumerate(items_s4):
        top_y = Inches(1.8 + i * 1.65)
        bar = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), top_y, Inches(0.2), Inches(1.4))
        bar.fill.solid()
        bar.fill.fore_color.rgb = icolor
        bar.line.color.rgb = icolor

        box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.3), top_y, Inches(11.0), Inches(1.4))
        box.fill.solid()
        box.fill.fore_color.rgb = WHITE
        box.line.color.rgb = CARD_BORDER
        box.line.width = Pt(1)

        tf = box.text_frame
        tf.word_wrap = True
        p_t = tf.paragraphs[0]
        p_t.text = ititle
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = NAVY

        p_d = tf.add_paragraph()
        p_d.text = idesc
        p_d.font.size = Pt(13)
        p_d.font.color.rgb = DARK_TEXT
        p_d.space_before = Pt(4)

    # ==========================================
    # SLIDE 5: EL DESAFÍO MULTI-MONEDA
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "La Realidad Comercial: Múltiples Monedas y Cajas", "EL RETO DEL DÍA A DÍA")

    challenges = [
        ("Efectivo Divisas (USD)", "Billetes en caja física que deben coincidir exactamente con los cobros de mostrador."),
        ("Bolívares (VES)", "Transferencias bancarias, puntos de venta y Pago Móvil sujetos a la tasa del día."),
        ("Dólar Digital (Zelle / USDT)", "Cuentas internacionales o billeteras con comisiones y tasas variables."),
        ("Operaciones de Traspaso", "Cuando se cambian divisas a bolívares para pagar proveedores o nómina.")
    ]
    for i, (ctitle, cdesc) in enumerate(challenges):
        col = i % 2
        row = i // 2
        left = Inches(1.0 + col * 5.8)
        top = Inches(1.8 + row * 2.5)

        card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(5.5), Inches(2.2))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1.5)

        bar = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(5.5), Inches(0.45))
        bar.fill.solid()
        bar.fill.fore_color.rgb = NAVY if i % 2 == 0 else DARK_BLUE
        bar.line.color.rgb = bar.fill.fore_color.rgb

        tf_b = bar.text_frame
        p_b = tf_b.paragraphs[0]
        p_b.text = ctitle
        p_b.font.size = Pt(13)
        p_b.font.bold = True
        p_b.font.color.rgb = GOLD if i % 2 == 0 else WHITE

        tf = card.text_frame
        tf.word_wrap = True
        p_pad = tf.paragraphs[0]
        p_pad.text = "\n" + cdesc
        p_pad.font.size = Pt(13)
        p_pad.font.color.rgb = DARK_TEXT

    # Bottom Callout
    bot_callout = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(6.1), Inches(11.3), Inches(0.85))
    bot_callout.fill.solid()
    bot_callout.fill.fore_color.rgb = GOLD
    bot_callout.line.color.rgb = GOLD
    tf_call = bot_callout.text_frame
    p_call = tf_call.paragraphs[0]
    p_call.text = "Conclusión: La gestión convencional en hojas dispersas genera retrasos y descuadres. El control inteligente los resuelve."
    p_call.font.size = Pt(13)
    p_call.font.bold = True
    p_call.font.color.rgb = NAVY
    p_call.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 6: LA SOLUCIÓN - SISTEMA TEV
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "La Solución: Sistema Integral de Tesorería TEV", "INNOVACIÓN Y AUTOMATIZACIÓN")

    features_s6 = [
        ("Nube 24/7", "Disponible desde cualquier computador de oficina, tablet o celular de los directivos."),
        ("Multi-Moneda Nativo", "Lleva saldos independientes en USD, VES y USDT con conversión transparente."),
        ("5 Operaciones Rápidas", "Tarjetas visuales e intuitivas para registrar gastos, compras, cobros y ventas."),
        ("Calculadora de Traspasos", "Registra salidas y entradas con la tasa pactada y detecta diferenciales cambiarios."),
        ("Matriz 1 al 31", "Visualización tradicional del mes completo, organizada día por día con sus saldos."),
        ("Reporte Ejecutivo 1-Click", "Ficha gerencial lista para imprimir o enviar en PDF a la junta directiva.")
    ]
    for i, (ftitle, fdesc) in enumerate(features_s6):
        col = i % 3
        row = i // 3
        left = Inches(1.0 + col * 3.9)
        top = Inches(1.8 + row * 2.5)

        card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(3.6), Inches(2.2))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1)

        tag = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left + Inches(0.3), top + Inches(0.2), Inches(3.0), Inches(0.45))
        tag.fill.solid()
        tag.fill.fore_color.rgb = NAVY
        tag.line.color.rgb = NAVY
        tf_tag = tag.text_frame
        p_t = tf_tag.paragraphs[0]
        p_t.text = ftitle
        p_t.font.size = Pt(12)
        p_t.font.bold = True
        p_t.font.color.rgb = GOLD
        p_t.alignment = PP_ALIGN.CENTER

        tb = s6.shapes.add_textbox(left + Inches(0.2), top + Inches(0.75), Inches(3.2), Inches(1.3))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = fdesc
        p.font.size = Pt(12)
        p.font.color.rgb = DARK_TEXT

    # ==========================================
    # SLIDE 7: ROLES Y FLUJO DE USO
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "¿Cómo se Utiliza? (Diseñado para Cada Rol)", "MANUAL OPERATIVO")

    roles = [
        ("ROL: CAJERA", DARK_BLUE, [
            "1. Abre el sistema en el navegador al iniciar turno.",
            "2. En 'Cierre Venta Diaria' registra el dinero cobrado del día.",
            "3. En 'Cobro CxC' registra abonos de clientes a crédito.",
            "4. Cuadra su caja en 3 minutos antes de entregar el turno."
        ]),
        ("ROL: ADMINISTRADORA", NAVY, [
            "1. Registra 'Gastos Operativos' y facturas de servicios.",
            "2. Registra 'Pago a Proveedores' con su soporte.",
            "3. Usa 'Cambio de Moneda' al convertir divisas a bolívares.",
            "4. Puede importar el archivo Excel mensual con 1 clic."
        ]),
        ("ROL: DIRECTIVA / DUEÑOS", GOLD, [
            "1. Ingresan desde su laptop o celular a cualquier hora.",
            "2. Observan el Saldo Consolidado Disponible en tiempo real.",
            "3. Consultan la Matriz Día por Día (1 al 31).",
            "4. Imprimen el Reporte Ejecutivo One-Page para reuniones."
        ])
    ]
    for i, (rtitle, rcolor, rsteps) in enumerate(roles):
        left = Inches(1.0 + i * 3.9)
        top = Inches(1.8)

        card = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(3.6), Inches(4.9))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1.5)

        banner = s7.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(3.6), Inches(0.65))
        banner.fill.solid()
        banner.fill.fore_color.rgb = rcolor
        banner.line.color.rgb = rcolor
        tf_b = banner.text_frame
        p_b = tf_b.paragraphs[0]
        p_b.text = rtitle
        p_b.font.size = Pt(13)
        p_b.font.bold = True
        p_b.font.color.rgb = NAVY if rcolor == GOLD else WHITE
        p_b.alignment = PP_ALIGN.CENTER

        tb = s7.shapes.add_textbox(left + Inches(0.2), top + Inches(0.8), Inches(3.2), Inches(3.9))
        tf = tb.text_frame
        tf.word_wrap = True
        for step_idx, step in enumerate(rsteps):
            p = tf.paragraphs[0] if step_idx == 0 else tf.add_paragraph()
            p.text = step
            p.font.size = Pt(12)
            p.font.color.rgb = DARK_TEXT
            p.space_before = Pt(8)

    # ==========================================
    # SLIDE 8: CARGADOR DE EXCEL Y FLEXIBILIDAD
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Compatibilidad Total: Su Excel Tradicional Sigue Vivo", "FLEXIBILIDAD SIN FRICCIONES")

    box_ex = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.8), Inches(11.3), Inches(4.8))
    box_ex.fill.solid()
    box_ex.fill.fore_color.rgb = WHITE
    box_ex.line.color.rgb = CARD_BORDER

    tf_ex = box_ex.text_frame
    tf_ex.word_wrap = True
    p = tf_ex.paragraphs[0]
    p.text = "Cargador Inteligente de Gastos Mensuales (Arrastrar y Soltar)"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = NAVY

    p_body = tf_ex.add_paragraph()
    p_body.text = (
        "\nUno de los mayores temores al adoptar nueva tecnología es perder el trabajo ya hecho o tener que escribir todo de nuevo.\n\n"
        "El Sistema TEV incluye un Importador Universal de Archivos Excel:\n\n"
        "• Arrastre cualquier archivo de gastos (Junio, Julio, Agosto, o el mes en curso).\n"
        "• El sistema interpreta automáticamente las columnas de Fecha, Concepto, Monto y Rubro.\n"
        "• Los gastos se integran al Flujo de Caja y a la Matriz 1-31 en menos de 2 segundos.\n"
        "• Puede crear nuevas cajas de monedas (ej. 'Caja Chica Sucursal 2') o nuevos rubros presupuestarios cuando lo necesite."
    )
    p_body.font.size = Pt(14)
    p_body.font.color.rgb = DARK_TEXT

    # ==========================================
    # SLIDE 9: BENEFICIOS Y RETORNO
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "El Retorno: ¿Qué Gana Todo Eléctrico Valencia?", "VALOR DE LA INVERSIÓN")

    benefits = [
        ("Tiempo Ahorrado", "De 2 horas diarias cuadrando números a reportes instantáneos.", DARK_BLUE),
        ("Cero Pérdidas Ocultas", "Control milimétrico del efectivo y de las diferencias por cambio de divisas.", NAVY),
        ("Poder de Negociación", "Saber exactamente cuándo pagar a proveedores para obtener mejores descuentos.", GOLD),
        ("Tranquilidad Gerencial", "Los directivos toman decisiones sobre datos reales, no sobre suposiciones.", DARK_BLUE)
    ]
    for i, (btitle, bdesc, bcolor) in enumerate(benefits):
        left = Inches(1.0 + i * 2.9)
        top = Inches(2.0)

        card = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(2.7), Inches(4.4))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1.5)

        circle = s9.shapes.add_shape(MSO_SHAPE.OVAL, left + Inches(0.85), top + Inches(0.4), Inches(1.0), Inches(1.0))
        circle.fill.solid()
        circle.fill.fore_color.rgb = bcolor
        circle.line.color.rgb = bcolor
        tf_c = circle.text_frame
        p_c = tf_c.paragraphs[0]
        p_c.text = "✓"
        p_c.font.size = Pt(28)
        p_c.font.bold = True
        p_c.font.color.rgb = WHITE if bcolor != GOLD else NAVY
        p_c.alignment = PP_ALIGN.CENTER

        tb = s9.shapes.add_textbox(left + Inches(0.15), top + Inches(1.6), Inches(2.4), Inches(2.5))
        tf = tb.text_frame
        tf.word_wrap = True
        p_t = tf.paragraphs[0]
        p_t.text = btitle
        p_t.font.size = Pt(15)
        p_t.font.bold = True
        p_t.font.color.rgb = NAVY
        p_t.alignment = PP_ALIGN.CENTER

        p_d = tf.add_paragraph()
        p_d.text = "\n" + bdesc
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = DARK_TEXT
        p_d.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 10: CIERRE
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    bg10 = s10.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg10.fill.solid()
    bg10.fill.fore_color.rgb = NAVY
    bg10.line.color.rgb = NAVY

    # Decorative Gold Bar
    bar10 = s10.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(1.6), Inches(0.15), Inches(4.2))
    bar10.fill.solid()
    bar10.fill.fore_color.rgb = GOLD
    bar10.line.color.rgb = GOLD

    # Logo
    if os.path.exists(logo_path):
        s10.shapes.add_picture(logo_path, Inches(1.6), Inches(1.5), height=Inches(1.2))

    tbox10 = s10.shapes.add_textbox(Inches(1.6), Inches(2.9), Inches(10.5), Inches(3.2))
    tf10 = tbox10.text_frame
    tf10.word_wrap = True

    p1 = tf10.paragraphs[0]
    p1.text = "El Futuro de la Gestión Financiera en Sus Manos"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = WHITE

    p2 = tf10.add_paragraph()
    p2.text = "Sistema Disponible en Línea: https://tev-tesoreria.onrender.com"
    p2.font.size = Pt(20)
    p2.font.color.rgb = GOLD
    p2.space_before = Pt(14)

    p3 = tf10.add_paragraph()
    p3.text = "Preguntas, Comentarios e Inicio de Pruebas Piloto"
    p3.font.size = Pt(16)
    p3.font.color.rgb = RGBColor(203, 213, 225)
    p3.space_before = Pt(16)

    output_path = os.path.abspath("PRESENTACION_EJECUTIVA_TEV.pptx")
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    create_presentation()
